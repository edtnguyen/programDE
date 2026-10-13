#!/usr/bin/env python3
"""Build the final IP/GV K30, dt=2.0 PerturbNMF summary workbook.

This adapts PerturbNMF Stage 3B to calibrated CRT outputs. It deliberately
retains beta_clr and the empirical/skew p-value columns rather than relabeling
the CRT effect as ordinary expression log2FC.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import math
from pathlib import Path
import re
import sys
from typing import Iterable

import muon as mu
import numpy as np
import pandas as pd
import scipy.sparse as sp
from scipy.stats import norm
from statsmodels.stats.multitest import multipletests

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_RUN_DIR = Path(
    "/oak/stanford/groups/engreitz/Users/tri/10xPerturbSeq/"
    "analysis/250507_ePAH1_NovaSeq/CRISPR_count/"
    "IP_GV_unified_sceptre_cnmf_no_celltype/PerturbNMF/results/"
    "IP_GV_skcnmf_K30_dt2_0_n100"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create the final multi-sheet workbook for the IP/GV K30, "
            "density-threshold=2.0 PerturbNMF model."
        )
    )
    parser.add_argument(
        "--run-dir",
        type=Path,
        default=DEFAULT_RUN_DIR,
    )
    parser.add_argument("--k", type=int, default=30)
    parser.add_argument("--density-threshold", type=float, default=2.0)
    parser.add_argument("--crt-variant", default="rna_qc_guides")
    parser.add_argument("--top-genes", type=int, default=300)
    parser.add_argument("--summary-top-genes", type=int, default=30)
    parser.add_argument("--fdr", type=float, default=0.05)
    parser.add_argument("--specificity-fdr", type=float, default=0.10)
    parser.add_argument(
        "--control-target",
        action="append",
        default=["non-targeting"],
        help=(
            "Guide target label treated as a control. May be repeated. "
            "Default: non-targeting."
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
    )
    parser.add_argument(
        "--output-name",
        default="IP_GV_K30_dt2_PerturbNMF_summary.xlsx",
    )
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def threshold_tag(value: float) -> str:
    return str(float(value)).replace(".", "_")


def read_required_tsv(path: Path) -> pd.DataFrame:
    if not path.is_file() or path.stat().st_size == 0:
        raise FileNotFoundError(path)
    frame = pd.read_csv(path, sep="\t")
    if frame.empty:
        raise ValueError(f"Input table is empty: {path}")
    return frame


def read_optional_tsv(path: Path) -> pd.DataFrame | None:
    if not path.is_file() or path.stat().st_size == 0:
        return None
    frame = pd.read_csv(path, sep="\t")
    return frame if not frame.empty else None


def first_column(
    frame: pd.DataFrame,
    candidates: Iterable[str],
    *,
    required: bool = True,
) -> str | None:
    for candidate in candidates:
        if candidate in frame.columns:
            return candidate
    if required:
        raise KeyError(
            f"None of {list(candidates)} occurs in columns "
            f"{frame.columns.tolist()}."
        )
    return None


def as_str_program(values: pd.Series | pd.Index) -> pd.Series:
    series = pd.Series(values, copy=False).astype(str)
    return series.str.replace(r"\.0$", "", regex=True)


def normalize_program_column(
    frame: pd.DataFrame,
    *,
    allow_index: bool = True,
) -> pd.DataFrame:
    result = frame.copy()

    if "program_name" not in result.columns and allow_index:
        unnamed = [
            column
            for column in result.columns
            if str(column).lower().startswith("unnamed")
        ]
        if unnamed:
            result = result.rename(columns={unnamed[0]: "program_name"})
        elif result.index.name == "program_name":
            result = result.reset_index()

    if "program_name" in result.columns:
        result["program_name"] = as_str_program(
            result["program_name"]
        ).to_numpy()
    return result


def adjusted_p_column(frame: pd.DataFrame) -> str | None:
    return first_column(
        frame,
        [
            "Adjusted P-value",
            "adj_pval",
            "primary_q_BH",
            "q_skew_BH",
            "q_empirical_BH",
            "FDR q-val",
        ],
        required=False,
    )


def truncate_excel_string(value: object, limit: int = 30_000) -> object:
    if isinstance(value, str) and len(value) > limit:
        return value[: limit - 15] + " ...[truncated]"
    return value


def excel_safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if result[column].dtype == object:
            result[column] = result[column].map(truncate_excel_string)
    return result


def resolve_guide_metadata(mdata):
    rna = mdata["rna"]
    programs = mdata["cNMF"]

    assignment = None
    for modality in [programs, rna]:
        for key in ["guide_assignment", "guide_assignments"]:
            if key in modality.obsm:
                assignment = modality.obsm[key]
                break
        if assignment is not None:
            break

    if assignment is None:
        raise KeyError("No guide-assignment matrix found in the H5MU.")

    if sp.issparse(assignment):
        assignment = assignment.tocsr().copy()
    else:
        assignment = sp.csr_matrix(np.asarray(assignment))
    assignment.eliminate_zeros()
    assignment.data = np.ones(assignment.nnz, dtype=np.int8)

    guide_names = None
    guide_targets = None
    for modality in [programs, rna]:
        if guide_names is None and "guide_names" in modality.uns:
            guide_names = np.asarray(
                modality.uns["guide_names"]
            ).astype(str)
        if guide_targets is None and "guide_targets" in modality.uns:
            guide_targets = np.asarray(
                modality.uns["guide_targets"]
            ).astype(str)

    if guide_names is None or guide_targets is None:
        raise KeyError("guide_names or guide_targets are absent.")
    if assignment.shape != (rna.n_obs, len(guide_names)):
        raise ValueError(
            f"Guide assignment has shape {assignment.shape}; expected "
            f"{(rna.n_obs, len(guide_names))}."
        )
    if len(guide_targets) != len(guide_names):
        raise ValueError("guide_names and guide_targets lengths differ.")

    return assignment, guide_names, guide_targets


def resolve_sample_series(mdata) -> pd.Series:
    rna = mdata["rna"]
    programs = mdata["cNMF"]

    if "sample" in programs.obs:
        values = programs.obs["sample"]
    elif "sample" in rna.obs:
        values = rna.obs.loc[programs.obs_names, "sample"]
    elif "sample_id" in rna.obs:
        values = rna.obs.loc[programs.obs_names, "sample_id"]
    else:
        raise KeyError("No sample or sample_id column found.")

    return pd.Series(
        values.astype(str).to_numpy(),
        index=programs.obs_names,
        name="sample",
    )


def gene_symbol_map(rna) -> dict[str, str]:
    if "symbol" in rna.var:
        symbols = rna.var["symbol"].astype(str)
    else:
        symbols = pd.Series(
            rna.var_names.astype(str),
            index=rna.var_names,
        )

    mapping: dict[str, str] = {}
    for var_name, symbol in zip(
        rna.var_names.astype(str),
        symbols.astype(str),
    ):
        mapping.setdefault(var_name, symbol)
        mapping.setdefault(symbol, symbol)
        mapping.setdefault(var_name.upper(), symbol)
        mapping.setdefault(symbol.upper(), symbol)
    return mapping


def compile_program_loadings(
    mdata,
    *,
    top_n: int,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rna = mdata["rna"]
    programs = mdata["cNMF"]

    if "loadings" not in programs.varm:
        raise KeyError("cNMF.varm['loadings'] is missing.")
    if "var_names" not in programs.uns:
        raise KeyError("cNMF.uns['var_names'] is missing.")

    matrix = np.asarray(programs.varm["loadings"], dtype=float)
    source_genes = np.asarray(programs.uns["var_names"]).astype(str)
    program_names = programs.var_names.astype(str)

    if matrix.shape != (programs.n_vars, len(source_genes)):
        raise ValueError(
            f"Loadings shape {matrix.shape} does not match "
            f"({programs.n_vars}, {len(source_genes)})."
        )

    mapping = gene_symbol_map(rna)
    display_genes = np.asarray(
        [
            mapping.get(gene, mapping.get(gene.upper(), gene))
            for gene in source_genes
        ],
        dtype=object,
    )

    all_loadings = pd.DataFrame(
        matrix,
        index=program_names,
        columns=display_genes,
    )
    all_loadings.index.name = "program_name"

    n = min(top_n, all_loadings.shape[1])
    long_rows = []
    flat_rows = []

    for program, values in all_loadings.iterrows():
        order = np.argsort(-values.to_numpy())[:n]
        genes = values.index.to_numpy()[order]
        scores = values.to_numpy()[order]

        flat = {"program_name": str(program)}
        for rank, gene in enumerate(genes, start=1):
            flat[f"rank_{rank}"] = str(gene)
        flat_rows.append(flat)

        for rank, (gene, score) in enumerate(
            zip(genes, scores),
            start=1,
        ):
            long_rows.append(
                {
                    "program_name": str(program),
                    "rank": rank,
                    "gene": str(gene),
                    "loading": float(score),
                }
            )

    flat_frame = pd.DataFrame(flat_rows)
    long_frame = pd.DataFrame(long_rows)
    return all_loadings, flat_frame, long_frame


def top_terms_for_program(
    frame: pd.DataFrame | None,
    program: str,
    *,
    n: int = 10,
    fdr: float = 0.05,
) -> str:
    if frame is None or "program_name" not in frame.columns:
        return ""

    subset = frame.loc[
        frame["program_name"].astype(str) == str(program)
    ].copy()
    if subset.empty:
        return ""

    p_column = adjusted_p_column(subset)
    if p_column is not None:
        subset[p_column] = pd.to_numeric(
            subset[p_column],
            errors="coerce",
        )
        significant = subset.loc[subset[p_column] < fdr]
        if not significant.empty:
            subset = significant
        subset = subset.sort_values(p_column)

    term_column = first_column(
        subset,
        ["Term", "term", "Gene_set", "Pathway"],
        required=False,
    )
    if term_column is None:
        return ""

    return "; ".join(
        subset[term_column]
        .astype(str)
        .drop_duplicates()
        .head(n)
        .tolist()
    )


def find_program_row(
    frame: pd.DataFrame | None,
    program: str,
) -> pd.Series | None:
    if frame is None:
        return None

    work = normalize_program_column(frame)
    if "program_name" not in work.columns:
        return None

    subset = work.loc[
        work["program_name"].astype(str) == str(program)
    ]
    return None if subset.empty else subset.iloc[0]


def categorical_summary(
    association: pd.DataFrame | None,
    program: str,
) -> dict[str, object]:
    row = find_program_row(association, program)
    if row is None:
        return {}

    output = {}
    for column, value in row.items():
        if column == "program_name":
            continue
        output[f"mouse_association_{column}"] = value
    return output


def explained_variance_value(
    frame: pd.DataFrame | None,
    program: str,
) -> float:
    row = find_program_row(frame, program)
    if row is None:
        return np.nan

    value_column = first_column(
        row.to_frame().T,
        [
            "VarianceExplained",
            "variance_explained",
            "Explained_Variance",
            "explained_variance",
        ],
        required=False,
    )
    if value_column is None:
        numeric = pd.to_numeric(row, errors="coerce")
        numeric = numeric.drop(labels=["program_name"], errors="ignore")
        return float(numeric.iloc[0]) if len(numeric) else np.nan
    return float(pd.to_numeric(row[value_column], errors="coerce"))


def ensure_crt_columns(frame: pd.DataFrame, condition: str) -> pd.DataFrame:
    result = normalize_program_column(frame)
    required = {
        "target_name",
        "program_name",
        "beta_clr",
        "p_empirical",
        "primary_pvalue",
        "primary_q_BH",
    }
    missing = required.difference(result.columns)
    if missing:
        raise KeyError(
            f"CRT table for {condition} lacks columns: {sorted(missing)}"
        )

    result["target_name"] = result["target_name"].astype(str)
    result["condition"] = condition

    for column in [
        "beta_clr",
        "p_empirical",
        "p_skew",
        "q_empirical_BH",
        "q_skew_BH",
        "primary_pvalue",
        "primary_q_BH",
        "approx_relative_log2fc",
    ]:
        if column in result:
            result[column] = pd.to_numeric(
                result[column],
                errors="coerce",
            )

    return result


def compute_cross_mouse(
    ip: pd.DataFrame,
    gv: pd.DataFrame,
) -> pd.DataFrame:
    key = ["target_name", "program_name"]

    left = ip[
        key + ["beta_clr", "primary_pvalue", "primary_q_BH"]
    ].rename(
        columns={
            "beta_clr": "beta_IP_GFP",
            "primary_pvalue": "p_IP_GFP",
            "primary_q_BH": "q_IP_GFP",
        }
    )
    right = gv[
        key + ["beta_clr", "primary_pvalue", "primary_q_BH"]
    ].rename(
        columns={
            "beta_clr": "beta_GV_GFP",
            "primary_pvalue": "p_GV_GFP",
            "primary_q_BH": "q_GV_GFP",
        }
    )

    merged = left.merge(right, on=key, how="inner")
    merged["same_sign"] = (
        np.sign(merged["beta_IP_GFP"])
        == np.sign(merged["beta_GV_GFP"])
    )

    p_ip = merged["p_IP_GFP"].clip(1e-300, 1.0)
    p_gv = merged["p_GV_GFP"].clip(1e-300, 1.0)

    z_ip = np.sign(merged["beta_IP_GFP"]) * norm.isf(p_ip / 2)
    z_gv = np.sign(merged["beta_GV_GFP"]) * norm.isf(p_gv / 2)
    z_combined = (z_ip + z_gv) / math.sqrt(2)

    merged["signed_stouffer_z"] = z_combined
    merged["signed_stouffer_p"] = (
        2 * norm.sf(np.abs(z_combined))
    )
    merged["signed_stouffer_q_BH"] = multipletests(
        merged["signed_stouffer_p"],
        method="fdr_bh",
    )[1]
    return merged


def summarize_crt_for_program(
    frame: pd.DataFrame,
    program: str,
    *,
    condition: str,
    fdr: float,
    specificity_fdr: float,
) -> dict[str, object]:
    subset = frame.loc[
        frame["program_name"].astype(str) == str(program)
    ].copy()
    if subset.empty:
        return {}

    significant = subset.loc[subset["primary_q_BH"] < fdr]
    positive = significant.loc[significant["beta_clr"] > 0].copy()
    negative = significant.loc[significant["beta_clr"] < 0].copy()

    strongest = subset.loc[
        subset["primary_q_BH"] < specificity_fdr
    ].copy()
    strongest["abs_beta"] = strongest["beta_clr"].abs()
    strongest = strongest.sort_values(
        ["abs_beta", "primary_q_BH"],
        ascending=[False, True],
    ).head(10)

    prefix = condition
    return {
        f"{prefix}_n_significant_targets_FDR{fdr:g}": int(
            len(significant)
        ),
        f"{prefix}_positive_targets": "; ".join(
            positive.sort_values(
                "beta_clr",
                ascending=False,
            )["target_name"].astype(str).tolist()
        ),
        f"{prefix}_negative_targets": "; ".join(
            negative.sort_values(
                "beta_clr",
                ascending=True,
            )["target_name"].astype(str).tolist()
        ),
        f"{prefix}_top_targets_q<{specificity_fdr:g}": "; ".join(
            (
                strongest["target_name"].astype(str)
                + " (βclr="
                + strongest["beta_clr"].map(lambda x: f"{x:.3g}")
                + ", q="
                + strongest["primary_q_BH"].map(
                    lambda x: f"{x:.3g}"
                )
                + ")"
            ).tolist()
        ),
    }


def summarize_cross_mouse_for_program(
    frame: pd.DataFrame,
    program: str,
    *,
    fdr: float,
) -> dict[str, object]:
    subset = frame.loc[
        frame["program_name"].astype(str) == str(program)
    ].copy()
    if subset.empty:
        return {}

    replicated = subset.loc[
        subset["same_sign"]
        & (subset["signed_stouffer_q_BH"] < fdr)
    ].sort_values("signed_stouffer_q_BH")

    return {
        "cross_mouse_n_same_sign_FDR": int(len(replicated)),
        "cross_mouse_same_sign_targets": "; ".join(
            replicated["target_name"].astype(str).tolist()
        ),
        "cross_mouse_top_targets": "; ".join(
            (
                replicated["target_name"].astype(str)
                + " (q="
                + replicated["signed_stouffer_q_BH"].map(
                    lambda x: f"{x:.3g}"
                )
                + ")"
            ).head(10).tolist()
        ),
    }


def calculate_usage_summaries(mdata) -> pd.DataFrame:
    programs = mdata["cNMF"]
    sample = resolve_sample_series(mdata)

    usage = np.asarray(programs.X, dtype=float)
    if usage.shape != (programs.n_obs, programs.n_vars):
        raise ValueError("Unexpected cNMF usage shape.")

    rows = []
    global_means = usage.mean(axis=0)

    for program_index, program in enumerate(
        programs.var_names.astype(str)
    ):
        row = {"program_name": program}

        for condition in sorted(sample.unique()):
            mask = sample.to_numpy() == condition
            values = usage[mask, program_index]
            row[f"mean_usage_{condition}"] = float(values.mean())
            row[f"median_usage_{condition}"] = float(
                np.median(values)
            )
            row[
                f"fraction_above_global_mean_{condition}"
            ] = float(np.mean(values > global_means[program_index]))

        rows.append(row)

    return pd.DataFrame(rows)


def expression_vector_cp10k(
    rna,
    *,
    gene_index: int,
) -> np.ndarray:
    totals = np.asarray(rna.X.sum(axis=1)).ravel().astype(float)
    totals[totals <= 0] = np.nan

    column = rna.X[:, gene_index]
    if sp.issparse(column):
        values = column.toarray().ravel().astype(float)
    else:
        values = np.asarray(column).ravel().astype(float)

    return values / totals * 10_000.0


def target_gene_index(rna, target: str) -> int | None:
    candidates = {
        str(target),
        str(target).upper(),
        str(target).lower(),
    }

    for index, var_name in enumerate(rna.var_names.astype(str)):
        if (
            var_name in candidates
            or var_name.upper() in candidates
            or var_name.lower() in candidates
        ):
            return index

    if "symbol" in rna.var:
        for index, symbol in enumerate(rna.var["symbol"].astype(str)):
            if (
                symbol in candidates
                or symbol.upper() in candidates
                or symbol.lower() in candidates
            ):
                return index
    return None


def build_target_summary(
    mdata,
    assignment,
    guide_names,
    guide_targets,
    crt_by_condition: dict[str, pd.DataFrame],
    cross_mouse: pd.DataFrame,
    *,
    control_targets: set[str],
    fdr: float,
    specificity_fdr: float,
) -> pd.DataFrame:
    rna = mdata["rna"]
    sample = resolve_sample_series(mdata)
    target_series = pd.Series(guide_targets, index=guide_names)

    all_targets = sorted(pd.unique(target_series.astype(str)))
    control_mask_by_guide = target_series.isin(control_targets).to_numpy()
    control_cells = np.asarray(
        assignment[:, control_mask_by_guide].sum(axis=1)
    ).ravel() > 0

    rows = []

    for target in all_targets:
        guide_mask = (target_series == target).to_numpy()
        target_cells = np.asarray(
            assignment[:, guide_mask].sum(axis=1)
        ).ravel() > 0

        row: dict[str, object] = {
            "target_name": target,
            "is_control": target in control_targets,
            "n_guides": int(guide_mask.sum()),
            "total_assigned_cells": int(target_cells.sum()),
        }

        gene_index = (
            None
            if target in control_targets
            else target_gene_index(rna, target)
        )
        cp10k = (
            None
            if gene_index is None
            else expression_vector_cp10k(
                rna,
                gene_index=gene_index,
            )
        )

        for condition in sorted(sample.unique()):
            condition_mask = sample.to_numpy() == condition
            condition_target = condition_mask & target_cells
            condition_control = (
                condition_mask
                & control_cells
                & ~target_cells
            )

            row[f"assigned_cells_{condition}"] = int(
                condition_target.sum()
            )
            row[f"control_cells_{condition}"] = int(
                condition_control.sum()
            )

            if cp10k is not None and condition_target.any():
                mean_target = float(
                    np.nanmean(cp10k[condition_target])
                )
            else:
                mean_target = np.nan

            if cp10k is not None and condition_control.any():
                mean_control = float(
                    np.nanmean(cp10k[condition_control])
                )
            else:
                mean_control = np.nan

            row[f"mean_target_CP10K_{condition}"] = mean_target
            row[f"mean_control_CP10K_{condition}"] = mean_control
            row[f"KD_efficiency_percent_{condition}"] = (
                (1.0 - mean_target / mean_control) * 100.0
                if np.isfinite(mean_target)
                and np.isfinite(mean_control)
                and mean_control > 0
                else np.nan
            )

            crt = crt_by_condition[condition]
            subset = crt.loc[crt["target_name"] == target].copy()
            significant = subset.loc[
                subset["primary_q_BH"] < fdr
            ]

            row[f"significant_programs_{condition}"] = "; ".join(
                significant.sort_values("primary_q_BH")[
                    "program_name"
                ].astype(str).tolist()
            )
            row[f"n_significant_programs_{condition}"] = int(
                len(significant)
            )

            strongest = subset.loc[
                subset["primary_q_BH"] < specificity_fdr
            ].copy()
            strongest["abs_beta"] = strongest["beta_clr"].abs()
            strongest = strongest.sort_values(
                ["abs_beta", "primary_q_BH"],
                ascending=[False, True],
            ).head(5)

            row[f"top_programs_{condition}"] = "; ".join(
                (
                    strongest["program_name"].astype(str)
                    + " (βclr="
                    + strongest["beta_clr"].map(
                        lambda x: f"{x:.3g}"
                    )
                    + ", q="
                    + strongest["primary_q_BH"].map(
                        lambda x: f"{x:.3g}"
                    )
                    + ")"
                ).tolist()
            )

        replicated = cross_mouse.loc[
            (cross_mouse["target_name"] == target)
            & cross_mouse["same_sign"]
            & (cross_mouse["signed_stouffer_q_BH"] < fdr)
        ].sort_values("signed_stouffer_q_BH")

        row["cross_mouse_same_sign_programs"] = "; ".join(
            replicated["program_name"].astype(str).tolist()
        )
        row["n_cross_mouse_same_sign_programs"] = int(
            len(replicated)
        )

        rows.append(row)

    return pd.DataFrame(rows)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for column_cells in worksheet.columns:
            letter = get_column_letter(column_cells[0].column)
            width = 10
            for cell in column_cells[:2000]:
                if cell.value is not None:
                    width = max(
                        width,
                        min(len(str(cell.value)) + 2, 55),
                    )
                if isinstance(cell.value, str) and len(cell.value) > 50:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )
            worksheet.column_dimensions[letter].width = width

        header_to_column = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for header, column in header_to_column.items():
            lowered = header.lower()
            if (
                "q_" in lowered
                or lowered.endswith("_q")
                or "primary_q" in lowered
                or "pvalue" in lowered
                or "p-value" in lowered
            ):
                letter = get_column_letter(column)
                worksheet.conditional_formatting.add(
                    f"{letter}2:{letter}{worksheet.max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="63BE7B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="F8696B",
                    ),
                )

    workbook.save(path)


def main() -> None:
    args = parse_args()

    if args.k <= 1:
        raise ValueError("K must be >1.")
    if args.top_genes < 1:
        raise ValueError("--top-genes must be positive.")
    if not (0 < args.fdr < 1):
        raise ValueError("--fdr must be between 0 and 1.")

    run_dir = args.run_dir.resolve()
    tag = threshold_tag(args.density_threshold)
    evaluation_dir = (
        run_dir
        / "Evaluation"
        / f"{args.k}_{tag}"
    )
    h5mu_path = (
        run_dir
        / "Inference"
        / "adata"
        / f"cNMF_{args.k}_{tag}.h5mu"
    )

    crt_root = (
        evaluation_dir
        / "CRT_notebook"
        / "discovery"
        / args.crt_variant
    )
    ip_path = crt_root / "IP_GFP" / "target_program_CRT.tsv"
    gv_path = crt_root / "GV_GFP" / "target_program_CRT.tsv"
    cross_path = crt_root / "cross_mouse_target_program_summary.tsv"

    go_path = (
        evaluation_dir
        / f"{args.k}_GO_term_enrichment.txt"
    )
    reactome_path = (
        evaluation_dir
        / f"{args.k}_geneset_enrichment.txt"
    )
    association_path = (
        evaluation_dir
        / f"{args.k}_categorical_association_results.txt"
    )
    variance_path = (
        evaluation_dir
        / f"{args.k}_Explained_Variance.txt"
    )
    variance_summary_path = (
        evaluation_dir
        / f"{args.k}_Explained_Variance_Summary.txt"
    )

    required = [
        h5mu_path,
        ip_path,
        gv_path,
        go_path,
        reactome_path,
        association_path,
        variance_path,
    ]
    missing = [
        str(path)
        for path in required
        if not path.is_file() or path.stat().st_size == 0
    ]
    if missing:
        raise FileNotFoundError(
            "Required summary inputs are missing:\n"
            + "\n".join(missing)
        )

    output_dir = (
        args.output_dir.resolve()
        if args.output_dir is not None
        else (
            run_dir
            / "Interpretation"
            / "Summary"
            / f"K{args.k}_dt{tag}"
        )
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / args.output_name

    if workbook_path.exists() and not args.overwrite:
        raise FileExistsError(
            f"{workbook_path} exists. Pass --overwrite to replace it."
        )

    print(f"Reading: {h5mu_path}")
    mdata = mu.read(h5mu_path)

    if "rna" not in mdata.mod or "cNMF" not in mdata.mod:
        raise KeyError("Expected RNA and cNMF modalities.")
    if mdata["cNMF"].n_vars != args.k:
        raise ValueError(
            f"Expected K={args.k}, found "
            f"{mdata['cNMF'].n_vars} programs."
        )

    program_names = mdata["cNMF"].var_names.astype(str).tolist()
    sample = resolve_sample_series(mdata)
    conditions = sorted(sample.unique().tolist())

    expected_conditions = {"IP_GFP", "GV_GFP"}
    if set(conditions) != expected_conditions:
        raise ValueError(
            f"Expected conditions {sorted(expected_conditions)}, "
            f"found {conditions}."
        )

    assignment, guide_names, guide_targets = resolve_guide_metadata(
        mdata
    )

    _, loadings_flat, loadings_long = compile_program_loadings(
        mdata,
        top_n=args.top_genes,
    )

    go = normalize_program_column(read_required_tsv(go_path))
    reactome = normalize_program_column(
        read_required_tsv(reactome_path)
    )
    association = normalize_program_column(
        read_required_tsv(association_path)
    )
    variance = normalize_program_column(
        read_required_tsv(variance_path)
    )
    variance_summary = read_optional_tsv(
        variance_summary_path
    )

    ip = ensure_crt_columns(
        read_required_tsv(ip_path),
        "IP_GFP",
    )
    gv = ensure_crt_columns(
        read_required_tsv(gv_path),
        "GV_GFP",
    )
    crt_by_condition = {
        "IP_GFP": ip,
        "GV_GFP": gv,
    }

    cross_mouse = read_optional_tsv(cross_path)
    if cross_mouse is None:
        print(
            "Cross-mouse CRT summary was absent; computing it "
            "from the two per-mouse tables."
        )
        cross_mouse = compute_cross_mouse(ip, gv)
        cross_mouse.to_csv(
            cross_path,
            sep="\t",
            index=False,
        )
    else:
        cross_mouse = normalize_program_column(cross_mouse)

    usage_summary = calculate_usage_summaries(mdata)

    flat_lookup = loadings_flat.set_index("program_name")
    summary_rows = []

    for program in program_names:
        row: dict[str, object] = {
            "program_name": program,
            "manual_annotation": "",
            "notes": "",
        }

        if program in flat_lookup.index:
            genes = [
                str(value)
                for value in flat_lookup.loc[program].dropna().tolist()
            ]
        else:
            genes = []

        row["top10_loaded_genes"] = "; ".join(genes[:10])
        row["top30_loaded_genes"] = "; ".join(
            genes[: args.summary_top_genes]
        )
        row["top_GO_terms"] = top_terms_for_program(
            go,
            program,
            n=10,
            fdr=args.fdr,
        )
        row["top_Reactome_terms"] = top_terms_for_program(
            reactome,
            program,
            n=10,
            fdr=args.fdr,
        )
        row["variance_explained"] = explained_variance_value(
            variance,
            program,
        )

        usage_row = usage_summary.loc[
            usage_summary["program_name"] == program
        ]
        if not usage_row.empty:
            row.update(
                usage_row.iloc[0]
                .drop(labels=["program_name"])
                .to_dict()
            )

        row.update(categorical_summary(association, program))

        row.update(
            summarize_crt_for_program(
                ip,
                program,
                condition="IP_GFP",
                fdr=args.fdr,
                specificity_fdr=args.specificity_fdr,
            )
        )
        row.update(
            summarize_crt_for_program(
                gv,
                program,
                condition="GV_GFP",
                fdr=args.fdr,
                specificity_fdr=args.specificity_fdr,
            )
        )
        row.update(
            summarize_cross_mouse_for_program(
                cross_mouse,
                program,
                fdr=args.fdr,
            )
        )
        summary_rows.append(row)

    program_summary = pd.DataFrame(summary_rows)

    target_summary = build_target_summary(
        mdata,
        assignment,
        guide_names,
        guide_targets,
        crt_by_condition,
        cross_mouse,
        control_targets=set(args.control_target),
        fdr=args.fdr,
        specificity_fdr=args.specificity_fdr,
    )

    guide_metadata = pd.DataFrame(
        {
            "guide_name": guide_names,
            "guide_target": guide_targets,
        }
    )
    guide_metadata["is_control"] = guide_metadata[
        "guide_target"
    ].isin(set(args.control_target))

    sample_summary = (
        sample.value_counts()
        .rename_axis("sample")
        .rename("n_cells")
        .reset_index()
    )

    input_manifest = pd.DataFrame(
        [
            {"input": "H5MU", "path": str(h5mu_path)},
            {"input": "GO", "path": str(go_path)},
            {"input": "Reactome", "path": str(reactome_path)},
            {
                "input": "Categorical association",
                "path": str(association_path),
            },
            {
                "input": "Explained variance",
                "path": str(variance_path),
            },
            {"input": "CRT IP_GFP", "path": str(ip_path)},
            {"input": "CRT GV_GFP", "path": str(gv_path)},
            {"input": "Cross-mouse CRT", "path": str(cross_path)},
        ]
    )

    metadata = pd.DataFrame(
        [
            {"field": "created_utc", "value": datetime.now(
                timezone.utc
            ).isoformat()},
            {"field": "run_dir", "value": str(run_dir)},
            {"field": "K", "value": args.k},
            {
                "field": "density_threshold",
                "value": args.density_threshold,
            },
            {"field": "CRT_variant", "value": args.crt_variant},
            {"field": "CRT_primary_effect", "value": "beta_clr"},
            {
                "field": "CRT_primary_q",
                "value": "primary_q_BH",
            },
            {"field": "FDR", "value": args.fdr},
            {
                "field": "specificity_FDR",
                "value": args.specificity_fdr,
            },
            {
                "field": "control_targets",
                "value": "; ".join(args.control_target),
            },
            {"field": "n_cells", "value": mdata["rna"].n_obs},
            {"field": "n_genes", "value": mdata["rna"].n_vars},
            {"field": "n_programs", "value": args.k},
            {"field": "n_guides", "value": len(guide_names)},
            {
                "field": "conditions",
                "value": "; ".join(conditions),
            },
        ]
    )

    crt_combined = pd.concat([ip, gv], ignore_index=True)

    sheets: list[tuple[str, pd.DataFrame]] = [
        ("Program_summary", program_summary),
        ("Target_summary", target_summary),
        ("Program_loadings_flat", loadings_flat),
        ("Program_loadings_long", loadings_long),
        ("GO", go),
        ("Reactome", reactome),
        ("CRT_all", crt_combined),
        ("CRT_IP_GFP", ip),
        ("CRT_GV_GFP", gv),
        ("Cross_mouse_CRT", cross_mouse),
        ("Categorical_assoc", association),
        ("Explained_variance", variance),
        ("Guide_metadata", guide_metadata),
        ("Sample_summary", sample_summary),
        ("Run_metadata", metadata),
        ("Input_manifest", input_manifest),
    ]
    if variance_summary is not None:
        sheets.insert(
            sheets.index(("Guide_metadata", guide_metadata)),
            ("Explained_var_summary", variance_summary),
        )

    temporary = workbook_path.with_suffix(".tmp.xlsx")
    if temporary.exists():
        temporary.unlink()

    with pd.ExcelWriter(
        temporary,
        engine="openpyxl",
    ) as writer:
        for sheet_name, frame in sheets:
            safe = excel_safe_frame(frame)
            if len(safe) > 1_048_575:
                raise ValueError(
                    f"Sheet {sheet_name} has too many rows for Excel: "
                    f"{len(safe):,}."
                )
            safe.to_excel(
                writer,
                sheet_name=sheet_name[:31],
                index=False,
            )

    format_workbook(temporary)

    if workbook_path.exists():
        workbook_path.unlink()
    temporary.replace(workbook_path)

    # Useful sidecars for version control and scripted follow-up.
    program_summary.to_csv(
        output_dir / "Program_summary.tsv",
        sep="\t",
        index=False,
    )
    target_summary.to_csv(
        output_dir / "Target_summary.tsv",
        sep="\t",
        index=False,
    )
    loadings_long.to_csv(
        output_dir / "Program_loadings_long.tsv.gz",
        sep="\t",
        index=False,
        compression="gzip",
    )
    crt_combined.to_csv(
        output_dir / "CRT_all.tsv.gz",
        sep="\t",
        index=False,
        compression="gzip",
    )

    manifest = {
        "workbook": str(workbook_path),
        "sheets": [
            {"name": name[:31], "rows": int(len(frame))}
            for name, frame in sheets
        ],
        "K": args.k,
        "density_threshold": args.density_threshold,
        "crt_variant": args.crt_variant,
        "fdr": args.fdr,
    }
    (output_dir / "workbook_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n"
    )

    # Reopen and perform final structural validation.
    check = load_workbook(workbook_path, read_only=True)
    observed_sheets = check.sheetnames
    check.close()

    required_sheets = {
        "Program_summary",
        "Target_summary",
        "Program_loadings_flat",
        "Program_loadings_long",
        "GO",
        "Reactome",
        "CRT_IP_GFP",
        "CRT_GV_GFP",
        "Cross_mouse_CRT",
        "Explained_variance",
    }
    missing_sheets = required_sheets.difference(observed_sheets)
    if missing_sheets:
        raise RuntimeError(
            f"Workbook is missing sheets: {sorted(missing_sheets)}"
        )
    if workbook_path.stat().st_size == 0:
        raise RuntimeError("Workbook was created but is empty.")

    print("\nWorkbook created successfully")
    print("Path:", workbook_path)
    print("Size:", f"{workbook_path.stat().st_size / 1024**2:.2f} MiB")
    print("Sheets:")
    for sheet in observed_sheets:
        print("  ", sheet)


if __name__ == "__main__":
    main()
